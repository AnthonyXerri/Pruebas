#!/usr/bin/env python3
"""
Añade al deck una diapositiva de contexto con el flujo completo de Dependencia:
CU1 (reconocimiento del grado y PIA) y CU2 (adjudicación de servicios comunitarios).
Formas nativas de PowerPoint, editables. Se inserta justo antes de la diapositiva de CU1.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER

GRANATE = RGBColor(0x96, 0x16, 0x28)
ROJO    = RGBColor(0xC0, 0x00, 0x00)
PIZARRA = RGBColor(0x3F, 0x46, 0x52)
GRIS    = RGBColor(0x76, 0x7C, 0x86)
BORDE   = RGBColor(0xE6, 0xE6, 0xE6)
FONDO   = RGBColor(0xF5, 0xF5, 0xF5)
NEGRO   = RGBColor(0x00, 0x00, 0x00)
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE  = 'Century Gothic'

MX  = 697841
W   = 12192000 - 2 * MX
SEP = 70000

CU1 = [('Grabación', 'Registro, alta del expediente y subsanación de documentación',
        'Subsanación documental'),
       ('Valoración', 'Citación, visita del valorador, dictamen, comisión y resolución del grado',
        'Citación manual y valoración en papel'),
       ('Resolución del PIA', 'Elaboración del plan de cuidados y resolución de la prestación',
        'Tramitación manual y secuencial')]
CU2 = [('Alta en lista de espera', 'Por servicio (ayuda a domicilio o teleasistencia) y lote',
        None),
       ('Gestión y contacto', 'Localización, visita al domicilio y firma del acuerdo',
        'Personas no localizadas'),
       ('Asignación de servicio', 'Por prelación de grado y capacidad del lote',
        'Prelación por grado'),
       ('Alta del servicio', 'Inicio efectivo de la prestación en el domicilio',
        None)]


def txt(s, x, y, cx, cy, partes, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(cx), Emu(cy))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    for t, b, c, sz in partes:
        r = p.add_run(); r.text = t
        r.font.name = FUENTE; r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = c
    return tb


def forma(s, tipo, x, y, cx, cy, relleno, linea=None):
    sh = s.shapes.add_shape(tipo, Emu(x), Emu(y), Emu(cx), Emu(cy))
    sh.fill.solid(); sh.fill.fore_color.rgb = relleno
    if linea is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = linea; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def banda(s, y, etiqueta, fases, color):
    """Rótulo del caso de uso + fases en flechas encadenadas + descripción y cuello debajo."""
    txt(s, MX, y, W, 260000, [(etiqueta, True, color, 11)])
    n = len(fases)
    fw = (W - SEP * (n - 1)) // n
    fy, fh = y + 330000, 560000
    for i, (nombre, desc, cuello) in enumerate(fases):
        x = MX + i * (fw + SEP)
        tipo = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sh = forma(s, tipo, x, fy, fw, fh, color)
        if tipo == MSO_SHAPE.CHEVRON:
            sh.adjustments[0] = 0.28
        else:
            sh.adjustments[0] = 0.28
        tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Emu(260000 if i else 150000); tf.margin_right = Emu(200000)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = nombre
        r.font.name = FUENTE; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLANCO

        # descripción de la fase
        txt(s, x + 120000, fy + fh + 110000, fw - 240000, 420000,
            [(desc, False, GRIS, 9)], PP_ALIGN.CENTER)
        # cuello de botella principal, si lo hay
        if cuello:
            txt(s, x + 120000, fy + fh + 560000, fw - 240000, 200000,
                [('● ', True, ROJO, 9), (cuello, True, ROJO, 9)], PP_ALIGN.CENTER)
    return fy + fh


prs = Presentation('PainPoints_Dependencia_Discapacidad.pptx')
layout = prs.slide_masters[0].slide_layouts[3]
s = prs.slides.add_slide(layout)
for ph in list(s.placeholders):
    if ph.placeholder_format.type != PP_PLACEHOLDER.SLIDE_NUMBER:
        ph._element.getparent().remove(ph._element)

txt(s, 518160, 360664, 10500000, 400000,
    [('Proceso de dependencia · de la solicitud al servicio', True, NEGRO, 20)])
txt(s, 521838, 900000, 10500000, 380000,
    [('Dos casos de uso sobre un mismo recorrido, ambos gestionados en SIDEMA', False, GRIS, 11)])

banda(s, 1380000, 'CU1 · Reconocimiento del grado y PIA', CU1, GRANATE)

# el traspaso entre casos de uso
py = 3230000
pill = forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX + W // 2 - 2600000, py, 5200000, 330000,
             FONDO, BORDE)
pill.adjustments[0] = 0.5
tf = pill.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_top = tf.margin_bottom = 0
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
for t, b, c in [('▼  PIA resuelto', True, NEGRO), ('  ·  derecho reconocido, falta prestarlo', False, GRIS)]:
    r = p.add_run(); r.text = t
    r.font.name = FUENTE; r.font.size = Pt(10); r.font.bold = b; r.font.color.rgb = c

banda(s, 3800000, 'CU2 · Adjudicación de servicios comunitarios (EELL)', CU2, PIZARRA)

txt(s, MX, 6200000, W, 250000,
    [('● ', True, ROJO, 9), ('Cuello de botella principal de cada fase', False, GRIS, 9),
     ('     ·     Sistemas: SIDEMA · SAP', False, GRIS, 9)])

# mover la nueva diapositiva delante de CU1 (posición 2, tras el índice)
lst = prs.slides._sldIdLst
ids = list(lst)
nuevo = ids[-1]
lst.remove(nuevo)
lst.insert(2, nuevo)

prs.save('PainPoints_Dependencia_Discapacidad.pptx')
print('OK ·', len(list(lst)), 'diapositivas')
