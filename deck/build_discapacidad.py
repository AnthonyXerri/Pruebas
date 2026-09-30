#!/usr/bin/env python3
"""
Añade el bloque de Discapacidad al deck de pain points, con el mismo
formato de tarjetas que el bloque de Dependencia.
Las cifras salen de los paneles de Celonis (PROD · Discapacidad),
vista ejecutiva —tablas por centro base— y vista directiva.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER

GRANATE = RGBColor(0x96, 0x16, 0x28)
ROJO    = RGBColor(0xC0, 0x00, 0x00)
GRIS    = RGBColor(0x76, 0x7C, 0x86)
BORDE   = RGBColor(0xE6, 0xE6, 0xE6)
NEGRO   = RGBColor(0x00, 0x00, 0x00)
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE  = 'Century Gothic'

MX, GAP, COLS = 697841, 260000, 3
CARD_W = (12192000 - 2 * MX - GAP * (COLS - 1)) // COLS
CARD_H, TOP, PAD = 1930000, 1700000, 250000


def txt(cont, x, y, cx, cy, partes, align=PP_ALIGN.LEFT, interlineado=None):
    tb = cont.shapes.add_textbox(Emu(x), Emu(y), Emu(cx), Emu(cy))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    par = tf.paragraphs[0]
    par.alignment = align
    if interlineado:
        par.line_spacing = interlineado
    for t, nb, col, sz in partes:
        r = par.add_run()
        r.text = t
        r.font.name = FUENTE
        r.font.size = Pt(sz)
        r.font.bold = nb
        r.font.color.rgb = col
    return tb


def tarjeta(s, x, y, n, titulo, cifra, unidad, contexto):
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Emu(x), Emu(y), Emu(CARD_W), Emu(CARD_H))
    card.adjustments[0] = 0.04
    card.fill.solid(); card.fill.fore_color.rgb = BLANCO
    card.line.color.rgb = BORDE; card.line.width = Pt(0.75)
    card.shadow.inherit = False
    card.text_frame.text = ''

    filete = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Emu(x + 150000), Emu(y), Emu(320000), Emu(52000))
    filete.fill.solid(); filete.fill.fore_color.rgb = GRANATE
    filete.line.fill.background(); filete.shadow.inherit = False
    filete.text_frame.text = ''

    txt(s, x + PAD, y + 215000, 400000, 200000, [(f'{n:02d}', True, GRIS, 10)])
    txt(s, x + PAD, y + 450000, CARD_W - 2 * PAD, 480000,
        [(titulo, True, NEGRO, 12)], interlineado=1.15)
    txt(s, x + PAD, y + 940000, CARD_W - 2 * PAD, 400000,
        [(cifra, True, GRANATE, 28)])
    txt(s, x + PAD, y + 1320000, CARD_W - 2 * PAD, 200000,
        [(unidad, False, GRIS, 9)])
    txt(s, x + PAD, y + 1520000, CARD_W - 2 * PAD, 340000,
        [(contexto, False, GRIS, 9)], interlineado=1.25)


def slide_dolores(prs, layout, titulo, subtitulo, dolores, pie=None):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type != PP_PLACEHOLDER.SLIDE_NUMBER:
            ph._element.getparent().remove(ph._element)
    txt(s, 518160, 360664, 9253637, 400000, [(titulo, True, NEGRO, 20)])
    txt(s, 521838, 900000, 10500000, 380000, [(subtitulo, False, GRIS, 11)])
    for i, (t, c, u, ctx) in enumerate(dolores):
        col, fila = i % COLS, i // COLS
        tarjeta(s, MX + col * (CARD_W + GAP), TOP + fila * (CARD_H + GAP),
                i + 1, t, c, u, ctx)
    if pie:
        txt(s, MX, TOP + 2 * (CARD_H + GAP) + 60000, 12192000 - 2 * MX, 300000,
            [(pie, False, GRIS, 9)])
    return s


# ════════════════════════════════════════════════════════════════════════
#  DISCAPACIDAD · Panel de Entradas — reconocimiento del grado
#  Cifras leídas de los paneles de Celonis (PROD), periodo 01/01–30/04/2026.
# ════════════════════════════════════════════════════════════════════════
DISCA = [
    ('Plazo de gestión de la solicitud',
     '77 d', 'de media por expediente',
     'Tiempo promedio de gestión que mide el panel para el periodo analizado.'),
    ('Un solo tramo se lleva la mitad',
     '44 d', 'el mayor de los cuatro tramos',
     'Más de la mitad del plazo total se consume en una única fase.'),
    ('Reclamaciones previas',
     '1.249', 'el 5,3 % de las entradas',
     'Expedientes que vuelven a entrar y generan retrabajo sobre lo ya resuelto.'),
    ('Carga desigual entre centros base',
     '1.712–4.609', 'solicitudes por centro',
     'El centro más cargado recibe casi el triple que el menos cargado.'),
    ('Calidad del dato por centro',
     '100 %', 'del Centro Base 10 como <18',
     'Un centro entero con el tramo de edad y el despistaje sin informar bien.'),
    ('Entradas sin despistaje asignado',
     '44,3 %', 'figuran como «Ninguno»',
     'Casi la mitad de las solicitudes entran sin tipo de despistaje.'),
]
PIE = ('Fuente: Celonis · Discapacidad (PROD) · Panel de Entradas, vistas ejecutiva y directiva · '
       'periodo 01/01/2026 – 30/04/2026 · Centros base 01 a 10.')

prs = Presentation('entrada.pptx')
layout = prs.slide_masters[0].slide_layouts[3]

# ── índice: corregir el salto de numeración ─────────────────────────────
ph = [x for x in prs.slides[1].placeholders
      if x.placeholder_format.type == PP_PLACEHOLDER.OBJECT][0]
for par in ph.text_frame.paragraphs:
    t = ''.join(r.text for r in par.runs)
    if t.strip().startswith('4. Próximos pasos'):
        par.runs[0].text = '3. Próximos pasos'
        for r in par.runs[1:]:
            r.text = ''

slide_dolores(prs, layout,
              'Pain points · Discapacidad',
              'Reconocimiento del grado de discapacidad · solicitudes de entrada en los centros base',
              DISCA, PIE)

# ── orden: portada, índice, CU1, CU2, Discapacidad, cierre ──────────────
lst = prs.slides._sldIdLst
ids = list(lst)
orden = [ids[0], ids[1], ids[2], ids[3], ids[5], ids[4]]
for e in ids:
    lst.remove(e)
for e in orden:
    lst.append(e)

prs.save('PainPoints_Dependencia_Discapacidad.pptx')
print('OK ·', len(list(lst)), 'diapositivas')
